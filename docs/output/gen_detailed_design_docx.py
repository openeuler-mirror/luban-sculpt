#!/usr/bin/env python3
"""从 docs/详细设计文档.md 生成麒麟软件风格 Word（python-docx，无 pandoc）。"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "详细设计文档.md"
OUT_PATH = Path(__file__).resolve().parent / "《Luban-Sculpt 详细设计文档-麒麟软件版式》.docx"


def _set_run_font(run, *, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size_pt: float | None = None):
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size_pt is not None:
        run.font.size = Pt(size_pt)


def _add_para(doc: Document, text: str, *, bold: bool = False, size: float = 12, align=None, east_asia="宋体"):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    _set_run_font(run, east_asia=east_asia, size_pt=size)
    return p


def _heading(doc: Document, text: str, level: int) -> None:
    # Strip markdown heading markers already stripped by caller
    p = doc.add_heading(text, level=min(level, 3))
    for run in p.runs:
        _set_run_font(run, east_asia="黑体", ascii_font="Arial", size_pt=16 - level)


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                _set_run_font(run, east_asia="黑体", size_pt=10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    _set_run_font(run, size_pt=9)


def _split_table_row(line: str) -> list[str]:
    parts = [c.strip() for c in line.strip().strip("|").split("|")]
    return parts


def _is_sep_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells if c)


def md_to_docx(md: str, doc: Document) -> None:
    lines = md.splitlines()
    i = 0
    in_code = False
    code_buf: list[str] = []

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                block = "\n".join(code_buf)
                p = doc.add_paragraph()
                run = p.add_run(block)
                _set_run_font(run, east_asia="仿宋", ascii_font="Consolas", size_pt=9)
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.startswith("# "):
            _heading(doc, line[2:].strip(), 1)
            i += 1
            continue
        if line.startswith("## "):
            _heading(doc, line[3:].strip(), 2)
            i += 1
            continue
        if line.startswith("### "):
            _heading(doc, line[4:].strip(), 3)
            i += 1
            continue
        if line.startswith("#### "):
            _heading(doc, line[5:].strip(), 3)
            i += 1
            continue

        # blockquote meta
        if line.startswith("> "):
            text = line[2:].strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            _add_para(doc, text, bold=False, size=10.5)
            i += 1
            continue

        # table
        if "|" in line and line.strip().startswith("|"):
            rows_raw: list[list[str]] = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip().startswith("|"):
                rows_raw.append(_split_table_row(lines[i]))
                i += 1
            if len(rows_raw) >= 2 and _is_sep_row(rows_raw[1]):
                headers = rows_raw[0]
                body = [r for r in rows_raw[2:] if not _is_sep_row(r)]
                # normalize col count
                ncols = len(headers)
                body = [r + [""] * (ncols - len(r)) for r in body]
                body = [r[:ncols] for r in body]
                _add_table(doc, headers, body)
            else:
                for r in rows_raw:
                    _add_para(doc, " | ".join(r), size=10)
            continue

        # hr
        if re.fullmatch(r"-{3,}", line.strip()):
            i += 1
            continue

        # list
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(3))
            text = re.sub(r"`([^`]+)`", r"\1", text)
            p = doc.add_paragraph(text, style="List Bullet" if m.group(2) in "-*" else "List Number")
            for run in p.runs:
                _set_run_font(run, size_pt=11)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # normal paragraph
        text = line.strip()
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = re.sub(r"`([^`]+)`", r"\1", text)
        text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
        _add_para(doc, text, size=11)
        i += 1


def build() -> Path:
    md = MD_PATH.read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # cover-ish title from first H1
    m = re.search(r"^# (.+)$", md, re.M)
    title = m.group(1).strip() if m else "Luban-Sculpt 详细设计文档"
    _add_para(doc, title, bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, east_asia="黑体")
    _add_para(doc, "文档编号：LS-2026-002　版本：v2.2　基线：luban_sculpt v0.1.0", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, "（由 docs/详细设计文档.md 自动生成；与源码冲突时以源码为准）", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()

    # skip first H1 in body to avoid duplicate
    body = re.sub(r"^# .+\n", "", md, count=1)
    md_to_docx(body, doc)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    out = build()
    print(f"wrote {out}")
